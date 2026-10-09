from __future__ import annotations

import io
import re
from typing import Any, Dict, List, Optional, Tuple
from urllib.parse import urlparse

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.auth import DEFAULT_REVIEW_ROLES, merged_role_definitions
from app.db import col


HEADERS = [
    "项目 / 系统",
    "标题",
    "负责人",
    "仓库 URL",
    "分支 / Tag / Commit",
    "用途说明",
    "环境",
    "暴露面",
    "业务关键性",
    "回滚方案",
    "回滚补充说明",
    "凭据名称或 ID",
    "扫描范围",
    "额外审核角色",
]

ENV_ALIASES = {
    "dev": "dev",
    "开发": "dev",
    "stage": "stage",
    "staging": "stage",
    "预发": "stage",
    "prod": "prod",
    "production": "prod",
    "生产": "prod",
    "other": "other",
    "其它": "other",
    "其他": "other",
}
EXPOSURE_ALIASES = {
    "internal": "internal",
    "内网": "internal",
    "public": "public",
    "公网": "public",
}
CRIT_ALIASES = {
    "low": "low",
    "低": "low",
    "medium": "medium",
    "中": "medium",
    "high": "high",
    "高": "high",
    "critical": "critical",
    "关键": "critical",
}
ROLLBACK_ALIASES = {
    "version_downgrade": "version_downgrade",
    "版本回退": "version_downgrade",
    "remove_dependency": "remove_dependency",
    "移除依赖": "remove_dependency",
    "replace_alternative": "replace_alternative",
    "替换为已准入替代组件": "replace_alternative",
    "替换替代组件": "replace_alternative",
    "feature_flag_off": "feature_flag_off",
    "关闭功能开关": "feature_flag_off",
    "isolate_degrade": "isolate_degrade",
    "隔离与降级": "isolate_degrade",
    "other": "other",
    "其它": "other",
    "其他": "other",
}


def _cell(v: Any) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _role_aliases() -> Dict[str, str]:
    out: Dict[str, str] = {}
    for role, meta in merged_role_definitions().items():
        out[role.lower()] = role
        name = str((meta or {}).get("name_zh") or "").strip()
        if name:
            out[name.lower()] = role
    return out


def _split_multi(raw: str) -> List[str]:
    return [x.strip() for x in re.split(r"[,，;；\n|/]", raw or "") if x.strip()]


def _looks_like_url(url: str) -> bool:
    try:
        parsed = urlparse(url)
    except Exception:
        return False
    return parsed.scheme in {"http", "https"} and bool(parsed.netloc)


def build_template_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "引入单"
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="183F89")
    for i, name in enumerate(HEADERS, 1):
        cell = ws.cell(1, i, name)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = 22
    ws.column_dimensions["D"].width = 42
    ws.column_dimensions["F"].width = 28
    example = [
        "my-service",
        "引入 debug 组件",
        "",
        "https://github.com/debug-js/debug.git",
        "master",
        "日志调试库",
        "dev",
        "internal",
        "low",
        "版本回退",
        "",
        "",
        "",
        "",
    ]
    for i, val in enumerate(example, 1):
        ws.cell(2, i, val)
    ws.data_validations.append(DataValidation(type="list", formula1='"dev,stage,prod,other"', allow_blank=True, sqref="G2:G5000"))
    ws.data_validations.append(DataValidation(type="list", formula1='"internal,public"', allow_blank=True, sqref="H2:H5000"))
    ws.data_validations.append(DataValidation(type="list", formula1='"low,medium,high,critical"', allow_blank=True, sqref="I2:I5000"))
    ws.data_validations.append(
        DataValidation(
            type="list",
            formula1='"版本回退,移除依赖,替换为已准入替代组件,关闭功能开关,隔离与降级"',
            allow_blank=True,
            sqref="J2:J5000",
        )
    )

    note = wb.create_sheet("填写说明")
    lines = [
        ["字段", "是否必填", "说明"],
        ["项目 / 系统", "是", "落项目标系统名称"],
        ["标题", "否", "空则使用项目名"],
        ["负责人", "否", "空则使用当前登录人"],
        ["仓库 URL", "是", "http(s) Git 地址"],
        ["分支 / Tag / Commit", "否", "默认 main"],
        ["用途说明", "是", "引入原因"],
        ["环境", "否", "dev / stage / prod / other，默认 dev"],
        ["暴露面", "否", "internal（内网）/ public（公网），默认 internal"],
        ["业务关键性", "否", "low / medium / high / critical，默认 low"],
        ["回滚方案", "否", "版本回退 / 移除依赖 / 替换为已准入替代组件 / 关闭功能开关 / 隔离与降级"],
        ["回滚补充说明", "否", "选择其它方案时填写"],
        ["凭据名称或 ID", "否", "私有仓库填写凭据库中的名称或 ID"],
        ["扫描范围", "否", "相对路径，逗号分隔；空则整仓"],
        ["额外审核角色", "否", "只能追加。默认已包含：项目经理、研发经理、产品经理、法务、安全"],
        ["", "", "确认后会创建引入单并立即提交扫描；并发扫描数由管理员在策略中配置。"],
    ]
    for r, row in enumerate(lines, 1):
        for c, val in enumerate(row, 1):
            cell = note.cell(r, c, val)
            if r == 1:
                cell.font = Font(bold=True)
        note.column_dimensions[get_column_letter(r and 1)].width = 22
    note.column_dimensions["A"].width = 22
    note.column_dimensions["B"].width = 12
    note.column_dimensions["C"].width = 70

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _lookup_credential(raw: str) -> Tuple[Optional[str], Optional[str]]:
    key = (raw or "").strip()
    if not key:
        return None, None
    creds = col("credentials")
    doc = creds.find_one({"$or": [{"credential_id": key}, {"name": key}]}, {"credential_id": 1, "name": 1})
    if not doc:
        return None, "凭据未匹配到凭据库，请改成已有名称/ID 或清空"
    return str(doc.get("credential_id")), None


def parse_workbook(data: bytes) -> Dict[str, Any]:
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as e:
        raise ValueError("无法解析 Excel，请使用下载的模板：%s" % e)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Excel 为空")
    header = [_cell(x) for x in rows[0]]
    index = {name: i for i, name in enumerate(header)}
    missing = [h for h in ("项目 / 系统", "仓库 URL", "用途说明") if h not in index]
    if missing:
        raise ValueError("模板缺少列：%s" % "、".join(missing))

    role_alias = _role_aliases()
    items: List[dict] = []
    seen: Dict[str, int] = {}
    for excel_row, raw in enumerate(rows[1:], 2):
        if raw is None or all(_cell(x) == "" for x in raw):
            continue

        def col(name: str) -> str:
            i = index.get(name)
            if i is None or i >= len(raw):
                return ""
            return _cell(raw[i])

        errors: List[str] = []
        warnings: List[str] = []
        project = col("项目 / 系统")
        purpose = col("用途说明")
        repo_url = col("仓库 URL")
        if not project:
            errors.append("项目 / 系统不能为空")
        if not purpose:
            errors.append("用途说明不能为空")
        if not repo_url:
            errors.append("仓库 URL 不能为空")
        elif not _looks_like_url(repo_url):
            errors.append("仓库 URL 须为 http(s) 地址")

        env = ENV_ALIASES.get(col("环境").lower(), "")
        if col("环境") and not env:
            errors.append("环境只能是 dev/stage/prod/other")
        exposure = EXPOSURE_ALIASES.get(col("暴露面").lower(), "")
        if col("暴露面") and not exposure:
            errors.append("暴露面只能是 internal/public 或 内网/公网")
        crit = CRIT_ALIASES.get(col("业务关键性").lower(), "")
        if col("业务关键性") and not crit:
            errors.append("业务关键性只能是 low/medium/high/critical")
        rb = ROLLBACK_ALIASES.get(col("回滚方案").lower(), "")
        if col("回滚方案") and not rb:
            errors.append("回滚方案不在可选列表中")

        extra_roles = []
        for token in _split_multi(col("额外审核角色")):
            mapped = role_alias.get(token.lower())
            if not mapped:
                errors.append("未知审核角色：%s" % token)
                continue
            if mapped in DEFAULT_REVIEW_ROLES:
                continue
            extra_roles.append(mapped)

        cred_raw = col("凭据名称或 ID")
        credential_id, cred_warn = _lookup_credential(cred_raw)
        if cred_warn:
            warnings.append(cred_warn)

        ref = col("分支 / Tag / Commit") or "main"
        dup_key = "%s|%s" % (repo_url.rstrip("/").lower(), ref.lower())
        if repo_url and dup_key in seen:
            errors.append("与第 %s 行仓库+分支重复" % seen[dup_key])
        elif repo_url:
            seen[dup_key] = excel_row

        include_paths = _split_multi(col("扫描范围"))
        notes = col("回滚补充说明")
        rollback = rb or "version_downgrade"
        if rollback == "other" and notes:
            rollback_plan = "other:%s" % notes
        else:
            rollback_plan = rollback

        items.append(
            {
                "row": excel_row,
                "project": project,
                "title": col("标题") or project,
                "owner": col("负责人"),
                "purpose": purpose,
                "repo_url": repo_url,
                "ref": ref,
                "environment": env or "dev",
                "exposure": exposure or "internal",
                "business_criticality": crit or "low",
                "rollback_plan": rollback_plan,
                "credential_id": credential_id or "",
                "credential_input": cred_raw,
                "include_paths": include_paths,
                "extra_review_roles": extra_roles,
                "errors": errors,
                "warnings": warnings,
            }
        )

    return {
        "total": len(items),
        "error_count": sum(1 for x in items if x["errors"]),
        "warning_count": sum(1 for x in items if x["warnings"] and not x["errors"]),
        "items": items,
        "default_review_roles": list(DEFAULT_REVIEW_ROLES),
    }


def validate_item(item: dict, *, peer_keys: Optional[Dict[str, int]] = None) -> dict:
    out = dict(item)
    errors: List[str] = []
    warnings: List[str] = list(item.get("warnings") or [])
    project = _cell(item.get("project"))
    purpose = _cell(item.get("purpose"))
    repo_url = _cell(item.get("repo_url"))
    if not project:
        errors.append("项目 / 系统不能为空")
    if not purpose:
        errors.append("用途说明不能为空")
    if not repo_url:
        errors.append("仓库 URL 不能为空")
    elif not _looks_like_url(repo_url):
        errors.append("仓库 URL 须为 http(s) 地址")
    ref = _cell(item.get("ref")) or "main"
    env_raw = _cell(item.get("environment"))
    env = ENV_ALIASES.get(env_raw.lower(), "") if env_raw else "dev"
    if env_raw and not env:
        errors.append("环境只能是 dev/stage/prod/other")
    exp_raw = _cell(item.get("exposure"))
    exposure = EXPOSURE_ALIASES.get(exp_raw.lower(), "") if exp_raw else "internal"
    if exp_raw and not exposure:
        errors.append("暴露面只能是 internal/public 或 内网/公网")
    crit_raw = _cell(item.get("business_criticality"))
    crit = CRIT_ALIASES.get(crit_raw.lower(), "") if crit_raw else "low"
    if crit_raw and not crit:
        errors.append("业务关键性只能是 low/medium/high/critical")
    rb_raw = _cell(item.get("rollback_plan"))
    rollback = ROLLBACK_ALIASES.get(rb_raw.lower(), rb_raw) if rb_raw else "version_downgrade"
    if rb_raw and rollback.split(":", 1)[0] not in ROLLBACK_ALIASES and rollback.split(":", 1)[0] != "other":
        errors.append("回滚方案不在可选列表中")
    role_alias = _role_aliases()
    extra_roles: List[str] = []
    raw_roles = item.get("extra_review_roles") or []
    if isinstance(raw_roles, str):
        raw_roles = _split_multi(raw_roles)
    for token in raw_roles:
        mapped = role_alias.get(str(token).strip().lower())
        if not mapped:
            errors.append("未知审核角色：%s" % token)
            continue
        if mapped in DEFAULT_REVIEW_ROLES:
            continue
        extra_roles.append(mapped)
    include_paths = item.get("include_paths") or []
    if isinstance(include_paths, str):
        include_paths = _split_multi(include_paths)
    include_paths = [str(x).strip() for x in include_paths if str(x).strip()]
    if peer_keys is not None and repo_url:
        key = "%s|%s" % (repo_url.rstrip("/").lower(), ref.lower())
        prev = peer_keys.get(key)
        if prev and prev != item.get("row"):
            errors.append("与第 %s 行仓库+分支重复" % prev)
        else:
            peer_keys[key] = int(item.get("row") or 0)
    cred_raw = _cell(item.get("credential_input") or item.get("credential_id"))
    credential_id, cred_warn = _lookup_credential(cred_raw)
    if cred_warn:
        warnings = [w for w in warnings if "凭据" not in w] + [cred_warn]
    elif cred_raw:
        warnings = [w for w in warnings if "凭据" not in w]
    out.update(
        {
            "project": project,
            "title": _cell(item.get("title")) or project,
            "purpose": purpose,
            "repo_url": repo_url,
            "ref": ref,
            "environment": env or "dev",
            "exposure": exposure or "internal",
            "business_criticality": crit or "low",
            "rollback_plan": rollback or "version_downgrade",
            "include_paths": include_paths,
            "extra_review_roles": extra_roles,
            "credential_id": credential_id or "",
            "errors": errors,
            "warnings": warnings,
        }
    )
    return out
